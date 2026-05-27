"""
Build the enhanced presentation with:
- Cover page (IOM302, group members, Amazon orange)
- Table of Contents
- Conclusion slides per section
- Page numbers
- Acknowledgements + Q&A slides
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Colors ──────────────────────────────────────────────────────────────────
ORANGE   = RGBColor(0xFF, 0x99, 0x00)   # Amazon orange
NAVY     = RGBColor(0x16, 0x1E, 0x2D)   # Dark navy
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
LTGRAY   = RGBColor(0xF4, 0xF4, 0xF4)
DKGRAY   = RGBColor(0x33, 0x33, 0x33)
MDGRAY   = RGBColor(0x77, 0x77, 0x77)
LTORANGE = RGBColor(0xFF, 0xCC, 0x66)

# ── Helpers ──────────────────────────────────────────────────────────────────

def txt(slide, text, L, T, W, H,
        sz=16, bold=False, italic=False,
        color=None, align=PP_ALIGN.LEFT, wrap=True):
    tb = slide.shapes.add_textbox(L, T, W, H)
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size  = Pt(sz)
    r.font.bold  = bold
    r.font.italic = italic
    r.font.name  = 'Calibri'
    if color:
        r.font.color.rgb = color
    return tb


def rect(slide, L, T, W, H, fill, line=False):
    s = slide.shapes.add_shape(1, L, T, W, H)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(0.5)
    else:
        s.line.fill.background()
    return s


def move_to(prs, from_idx, to_idx):
    lst  = prs.slides._sldIdLst
    elem = lst[from_idx]
    lst.remove(elem)
    lst.insert(to_idx, elem)


def clear_placeholders(slide):
    for ph in list(slide.placeholders):
        sp = ph._element
        sp.getparent().remove(sp)


def pagenum(slide, W, H, num, light=False):
    color = RGBColor(0xAA, 0xAA, 0xAA) if not light else WHITE
    t = slide.shapes.add_textbox(W - Inches(1.1), H - Inches(0.45),
                                  Inches(0.9), Inches(0.38))
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = str(num)
    r.font.size  = Pt(11)
    r.font.name  = 'Calibri'
    r.font.color.rgb = color


def new_blank(prs):
    """Add a new blank slide and return it."""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    clear_placeholders(slide)
    return slide


# ── Load ─────────────────────────────────────────────────────────────────────
prs = Presentation('/home/user/clawdbot-space/presentation.pptx')
W, H = prs.slide_width, prs.slide_height


# ════════════════════════════════════════════════════════════════════════════
# COVER SLIDE  (reuse existing empty slide 0)
# ════════════════════════════════════════════════════════════════════════════
cover = prs.slides[0]

# Full background
rect(cover, 0, 0, W, H, NAVY)

# Orange top bar
rect(cover, 0, 0, W, Inches(0.65), ORANGE)
# Orange bottom bar
rect(cover, 0, H - Inches(0.65), W, Inches(0.65), ORANGE)

# Thin left accent line
rect(cover, Inches(0.55), Inches(0.65), Inches(0.08), H - Inches(1.3), ORANGE)

# Course code
txt(cover, 'IOM302  |  E-Commerce Business Strategy',
    Inches(0.8), Inches(0.8), W - Inches(1.6), Inches(0.5),
    sz=14, bold=True, color=LTORANGE)

# Main title
txt(cover, 'E-Commerce Business', Inches(0.8), Inches(1.4), W - Inches(1.6), Inches(0.9),
    sz=42, bold=True, color=WHITE)
txt(cover, 'Strategies Analysis', Inches(0.8), Inches(2.2), W - Inches(1.6), Inches(0.9),
    sz=42, bold=True, color=ORANGE)

# Divider
rect(cover, Inches(0.8), Inches(3.12), Inches(4.5), Inches(0.055), ORANGE)

# Group number
txt(cover, 'Group Number:  占位 (Placeholder)',
    Inches(0.8), Inches(3.25), Inches(5), Inches(0.42),
    sz=13, color=MDGRAY)

# Members label
txt(cover, 'Group Members', Inches(0.8), Inches(3.75), Inches(3.5), Inches(0.38),
    sz=13, bold=True, color=ORANGE)

# Members list
members = [
    ('Yiwen Hao',      '2253659',  'Alibaba'),
    ('Kaiwen Jiang',   '2251793',  'Amazon'),
    ('Xin He',         '2145653',  'Warby Parker'),
    ('Xuan Wu',        '2254778',  'Tesla'),
    ('Youtai Wang',    '2251996',  'Zara'),
    ('Zezheng Liu',    '2143890',  ''),
]
yy = Inches(4.18)
for name, sid, company in members:
    txt(cover, f'{name:<22}  {sid}', Inches(0.9), yy, Inches(5.2), Inches(0.33),
        sz=12, color=WHITE)
    yy += Inches(0.33)

# Right panel – company list
rect(cover, W - Inches(3.3), Inches(1.3), Inches(2.9), Inches(4.2), RGBColor(0x1F, 0x2B, 0x3E))
txt(cover, 'Case Studies', W - Inches(3.1), Inches(1.5), Inches(2.5), Inches(0.4),
    sz=12, bold=True, color=ORANGE)
rect(cover, W - Inches(3.1), Inches(1.95), Inches(2.5), Inches(0.04), ORANGE)

co_list = ['Alibaba', 'Amazon', 'Warby Parker', 'Tesla', 'Zara']
yy = Inches(2.1)
for co in co_list:
    txt(cover, f'▸  {co}', W - Inches(3.1), yy, Inches(2.6), Inches(0.38),
        sz=13, color=WHITE)
    yy += Inches(0.44)


# ════════════════════════════════════════════════════════════════════════════
# TABLE OF CONTENTS  (new slide, inserted at index 1)
# ════════════════════════════════════════════════════════════════════════════
# After adding TOC → it'll be at the end; we'll move it to index 1.
#
# Final page numbers for the TOC (1-indexed):
#   Alibaba       → 7
#   Amazon        → 14
#   Warby Parker  → 18
#   Tesla         → 23
#   Zara          → 28
#   Ack & Q&A     → 35
toc_entries = [
    ('01', 'Alibaba',               7),
    ('02', 'Amazon',               14),
    ('03', 'Warby Parker',         18),
    ('04', 'Tesla',                23),
    ('05', 'Zara',                 28),
    ('06', 'Acknowledgements & Q&A', 35),
]

toc = new_blank(prs)
rect(toc, 0, 0, W, H, WHITE)
# Top navy header
rect(toc, 0, 0, W, Inches(1.15), NAVY)
rect(toc, 0, Inches(1.15), W, Inches(0.07), ORANGE)
# Left accent
rect(toc, 0, 0, Inches(0.08), H, ORANGE)

txt(toc, 'Table of Contents',
    Inches(0.4), Inches(0.22), W - Inches(1), Inches(0.75),
    sz=34, bold=True, color=WHITE)

row_h = Inches(0.80)
start_y = Inches(1.38)
for i, (num, label, pg) in enumerate(toc_entries):
    yy = start_y + i * row_h
    # Alternating row background
    if i % 2 == 0:
        rect(toc, Inches(0.4), yy - Inches(0.06), W - Inches(0.55), row_h - Inches(0.04),
             RGBColor(0xFA, 0xFA, 0xFA))

    # Number badge
    rect(toc, Inches(0.5), yy + Inches(0.1), Inches(0.46), Inches(0.46), ORANGE)
    txt(toc, num, Inches(0.5), yy + Inches(0.08), Inches(0.46), Inches(0.46),
        sz=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Label
    txt(toc, label, Inches(1.15), yy + Inches(0.08), W - Inches(2.8), Inches(0.5),
        sz=19, color=DKGRAY)

    # Page number
    txt(toc, str(pg), W - Inches(1.5), yy + Inches(0.08), Inches(1.2), Inches(0.5),
        sz=17, bold=True, color=ORANGE, align=PP_ALIGN.RIGHT)

    # Separator
    if i < len(toc_entries) - 1:
        rect(toc, Inches(1.1), yy + row_h - Inches(0.1),
             W - Inches(1.6), Inches(0.012), RGBColor(0xE0, 0xE0, 0xE0))

# Footer
rect(toc, 0, H - Inches(0.42), W, Inches(0.42), NAVY)
txt(toc, 'IOM302  |  E-Commerce Business Strategies Analysis',
    Inches(0.5), H - Inches(0.38), W - Inches(1), Inches(0.35),
    sz=10, color=LTORANGE)

# Move TOC to index 1
move_to(prs, len(prs.slides) - 1, 1)

# ── State after cover + TOC insertion ─────────────────────────────────────
# idx: 0=Cover 1=TOC 2=img1 3=img2 4=img3 5=img4
#      6=Ali.title 7=Ali.strat 8=Ali.BM 9=Ali.risks 10=Ali.recs
#      11=divider 12=Amz.ovr 13=Amz.tbl 14=Amz.risks
#      15=WP.title 16=WP.strat 17=WP.strat2 18=WP.recs
#      19=Tesla.title 20=Tesla.strat 21=Tesla.risks 22=Tesla.sugg
#      23=Zara.title 24=Zara.strat 25=Zara.data 26=Zara.BM 27=Zara.risk
#      28=References
# Total: 29 slides


# ════════════════════════════════════════════════════════════════════════════
# CONCLUSION SLIDE FACTORY
# ════════════════════════════════════════════════════════════════════════════
def make_conclusion(prs, company, points):
    s = new_blank(prs)
    W, H = prs.slide_width, prs.slide_height

    # Background
    rect(s, 0, 0, W, H, WHITE)
    # Left accent bar
    rect(s, 0, 0, Inches(0.38), H, ORANGE)
    # Top header
    rect(s, 0, 0, W, Inches(1.05), NAVY)
    rect(s, 0, Inches(1.05), W, Inches(0.07), ORANGE)

    # Title
    txt(s, f'{company}  —  Section Conclusion',
        Inches(0.55), Inches(0.18), W - Inches(0.75), Inches(0.75),
        sz=26, bold=True, color=WHITE)

    # "Key Takeaways" label
    txt(s, 'KEY TAKEAWAYS',
        Inches(0.55), Inches(1.25), Inches(3.5), Inches(0.38),
        sz=11, bold=True, color=MDGRAY)

    rect(s, Inches(0.55), Inches(1.6), Inches(6), Inches(0.045), ORANGE)

    yy = Inches(1.75)
    for i, point in enumerate(points, 1):
        # Number badge
        rect(s, Inches(0.55), yy + Inches(0.05), Inches(0.37), Inches(0.37), ORANGE)
        txt(s, str(i), Inches(0.55), yy + Inches(0.03), Inches(0.37), Inches(0.37),
            sz=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        # Point
        txt(s, point, Inches(1.05), yy, W - Inches(1.5), Inches(0.65),
            sz=15, color=DKGRAY)
        yy += Inches(0.88)

    # Footer
    rect(s, 0, H - Inches(0.42), W, Inches(0.42), NAVY)
    txt(s, 'IOM302  |  E-Commerce Business Strategies Analysis',
        Inches(0.5), H - Inches(0.38), W - Inches(1), Inches(0.35),
        sz=10, color=LTORANGE)

    return s


conclusions_data = [
    ('Alibaba', [
        '"User First, AI-Driven" strategy leverages AI personalisation and 88VIP membership to deepen customer intimacy.',
        'Platform ecosystem model generates 71.7 % of revenue from customer management services rather than direct sales.',
        'Margin pressure from AI investment and intensifying competition requires balanced efficiency and growth strategy.',
    ]),
    ('Amazon', [
        'Customer intimacy, long-tail marketplace, and FBA fulfilment form a self-reinforcing e-commerce flywheel.',
        'Hybrid revenue model (products, seller fees, ads, Prime) diversifies income and reduces reliance on single streams.',
        'Governance of data privacy, seller relationships, cost efficiency, and sustainability are key ongoing challenges.',
    ]),
    ('Warby Parker', [
        'DTC omnichannel strategy successfully disrupted the traditional eyewear duopoly with design-forward, affordable products.',
        'Home try-on innovation and physical retail expansion create a differentiated, high-engagement customer journey.',
        'Scalability of cost structure and defensibility of brand differentiation determine long-term competitive position.',
    ]),
    ('Tesla', [
        'DTC e-commerce model with online customisation eliminates dealer costs and enables direct customer data collection.',
        'Vertical integration (battery, software, manufacturing) underpins both cost leadership and product differentiation.',
        'Reputational risks and intensifying EV competition require resilient brand management and strategy diversification.',
    ]),
    ('Zara', [
        'Agile supply chain enables a 2-week design-to-shelf turnaround, creating trend-responsiveness and scarcity urgency.',
        'Omnichannel integration aligns online discovery with in-store fulfilment for a seamless premium retail experience.',
        'Environmental pressure from fast fashion and data privacy regulations demand proactive sustainability governance.',
    ]),
]

# Add all 5 conclusions to the end first, then move each into place.
for company, points in conclusions_data:
    make_conclusion(prs, company, points)

# Now: 34 slides (0-33)
# Indices 29-33 = Ali.c, Amz.c, WP.c, Tesla.c, Zara.c  (in that order)
#
# Insertion plan (process last→first for correct index tracking):
#   Zara.c  (33) → after Zara.risk  (27) → move to 28
#   Tesla.c (32) → after Tesla.sugg (22) → move to 23
#   WP.c    (31) → after WP.recs   (18) → move to 19
#   Amz.c   (30) → after Amz.risks (14) → move to 15
#   Ali.c   (29) → after Ali.recs  (10) → move to 11
#
# Because we process last→first, earlier insertions don't shift indices of
# slides that were already in place before this block.

# Process last-to-first: always take from index 33 (the tail of the list)
# because as each conclusion is moved earlier, the next one becomes the new tail.
move_to(prs, 33, 28)   # Zara.c   → after Zara.risk  (idx 27) → pos 28
move_to(prs, 33, 23)   # Tesla.c  → after Tesla.sugg (idx 22) → pos 23
move_to(prs, 33, 19)   # WP.c     → after WP.recs    (idx 18) → pos 19
move_to(prs, 33, 15)   # Amz.c    → after Amz.risks  (idx 14) → pos 15
move_to(prs, 33, 11)   # Ali.c    → after Ali.recs   (idx 10) → pos 11

# ── Final content layout ──────────────────────────────────────────────────
# idx 0=Cover, 1=TOC, 2-5=imgs
#  6=Ali.title  7=Ali.strat  8=Ali.BM  9=Ali.risks  10=Ali.recs
# 11=Ali.c
# 12=divider  13=Amz.ovr  14=Amz.tbl  15=Amz.risks
# 16=Amz.c
# 17=WP.title  18=WP.strat  19=WP.strat2  20=WP.recs
# 21=WP.c
# 22=Tesla.title  23=Tesla.strat  24=Tesla.risks  25=Tesla.sugg
# 26=Tesla.c
# 27=Zara.title  28=Zara.strat  29=Zara.data  30=Zara.BM  31=Zara.risk
# 32=Zara.c
# 33=References
# Total now: 34 slides


# ════════════════════════════════════════════════════════════════════════════
# ACKNOWLEDGEMENTS SLIDE  (index 34)
# ════════════════════════════════════════════════════════════════════════════
ack = new_blank(prs)
rect(ack, 0, 0, W, H, NAVY)
rect(ack, 0, 0, W, Inches(0.65), ORANGE)
rect(ack, 0, H - Inches(0.65), W, Inches(0.65), ORANGE)

txt(ack, 'Acknowledgements',
    Inches(0.7), Inches(0.9), W - Inches(1.4), Inches(0.8),
    sz=36, bold=True, color=WHITE)

rect(ack, Inches(0.7), Inches(1.75), Inches(4), Inches(0.06), ORANGE)

ack_text = (
    'We would like to express our sincere gratitude to our course instructor '
    'and teaching assistants for their valuable guidance throughout IOM302.\n\n'
    'We also thank the authors and publishers of the academic works cited in this '
    'presentation for their foundational contributions to the field of e-commerce strategy.\n\n'
    'This group project was a collaborative effort, and we are grateful for '
    'each team member\'s dedication and hard work.'
)
txt(ack, ack_text,
    Inches(0.7), Inches(2.0), W - Inches(1.4), Inches(3.2),
    sz=17, color=RGBColor(0xCC, 0xCC, 0xCC))

# Team credits at bottom
txt(ack, 'Group Members:',
    Inches(0.7), H - Inches(1.6), Inches(3), Inches(0.4),
    sz=12, bold=True, color=ORANGE)
members_str = '  ·  '.join([f'{m[0]} ({m[1]})' for m in members])
txt(ack, members_str,
    Inches(0.7), H - Inches(1.2), W - Inches(1.4), Inches(0.45),
    sz=11, color=RGBColor(0x99, 0x99, 0x99))


# ════════════════════════════════════════════════════════════════════════════
# Q&A SLIDE  (index 35)
# ════════════════════════════════════════════════════════════════════════════
qa = new_blank(prs)
rect(qa, 0, 0, W, H, NAVY)
rect(qa, 0, 0, W, Inches(0.65), ORANGE)
rect(qa, 0, H - Inches(0.65), W, Inches(0.65), ORANGE)

# Large Q&A text centred
txt(qa, 'Q & A',
    Inches(0), Inches(1.8), W, Inches(2.0),
    sz=80, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)

txt(qa, 'Questions & Answers',
    Inches(0.5), Inches(3.85), W - Inches(1.0), Inches(0.55),
    sz=22, color=WHITE, align=PP_ALIGN.CENTER)

rect(qa, Inches(3.2), Inches(4.5), W - Inches(6.4), Inches(0.055), ORANGE)

txt(qa, 'Thank you for listening!',
    Inches(0.5), Inches(4.7), W - Inches(1.0), Inches(0.5),
    sz=17, color=RGBColor(0xAA, 0xAA, 0xAA), align=PP_ALIGN.CENTER)

txt(qa, 'IOM302  |  E-Commerce Business Strategies Analysis',
    Inches(0.5), Inches(5.4), W - Inches(1.0), Inches(0.45),
    sz=13, color=LTORANGE, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════════════════════
# PAGE NUMBERS  (add to every slide)
# ════════════════════════════════════════════════════════════════════════════
# Dark slides (navy bg): white page numbers
# Light slides: gray page numbers
dark_indices = {0, 34, 35}  # Cover, Acknowledgements, Q&A

for idx, slide in enumerate(prs.slides):
    pg = idx + 1
    light = idx in dark_indices
    pagenum(slide, W, H, pg, light=light)


# ── Save ─────────────────────────────────────────────────────────────────────
out_path = '/home/user/clawdbot-space/presentation_final.pptx'
prs.save(out_path)
print(f'Saved → {out_path}')
print(f'Total slides: {len(prs.slides)}')
